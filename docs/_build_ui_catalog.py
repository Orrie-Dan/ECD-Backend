"""Capture ECD UI screens into a Word catalog. Credentials via env vars only."""

from __future__ import annotations

import atexit
import json
import os
import re
import time
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image
from playwright.sync_api import TimeoutError as PlaywrightTimeout
from playwright.sync_api import Error as PlaywrightError
from playwright.sync_api import sync_playwright

BASE = os.environ.get("ECD_UI_BASE", "http://localhost:5174")
OUT_DIR = Path(os.environ.get("ECD_UI_OUT", r"D:\Esri\ECD Backend\docs\ui-catalog"))
SHOT_DIR = OUT_DIR / "shots"
DOCX_PATH = Path(
    os.environ.get("ECD_UI_DOCX", r"D:\Esri\ECD Backend\docs\ECD-system-screens.docx")
)
SLICE_HEIGHT = 2200

CREDS = {
    "caretaker": (
        os.environ["ECD_USER_CARETAKER"],
        os.environ["ECD_PASS_CARETAKER"],
    ),
    "director": (
        os.environ["ECD_USER_DIRECTOR"],
        os.environ["ECD_PASS_DIRECTOR"],
    ),
    "district": (
        os.environ["ECD_USER_DISTRICT"],
        os.environ["ECD_PASS_DISTRICT"],
    ),
    "ncda": (
        os.environ["ECD_USER_NCDA"],
        os.environ["ECD_PASS_NCDA"],
    ),
}


MANIFEST = OUT_DIR / "manifest.json"


def load_catalog() -> list[dict]:
    if MANIFEST.exists():
        try:
            return json.loads(MANIFEST.read_text(encoding="utf-8"))
        except json.JSONDecodeError:
            return []
    return []


def save_catalog(catalog: list[dict]) -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    MANIFEST.write_text(json.dumps(catalog, indent=2), encoding="utf-8")


def already(catalog: list[dict], section: str, title: str) -> bool:
    return any(
        item.get("section") == section
        and item.get("title") == title
        and item.get("files")
        for item in catalog
    )


PATH_CLICK_TEXT = {
    "/caretaker/kwiyandikisha": "Andikisha Umwana",
    "/caretaker/imirire": "Imirire",
    "/caretaker/sted": "Gutahura ubumuga",
    "/caretaker/raporo": "Raporo",
    "/caretaker/igenamiterere": "Igenamiterere",
    "/caretaker/ibindi": "Ibindi",
    "/caretaker/ikigo": "Ikigo",
    "/caretaker/igitabo": "Igitabo cya ECD",
    "/caretaker/imicungire": "Imicungire y'ikigo",
    "/caretaker/abakoresha": "Abakoresha",
    "/caretaker/kwimura": "Kwimura",
    "/caretaker/abana": "Abana",
    "/caretaker/ubwitabire": "Ubwitabire",
    "/caretaker/imikurire": "Imikurire",
}


def slug(text: str) -> str:
    text = re.sub(r"[^a-zA-Z0-9]+", "-", text).strip("-").lower()
    return text[:80] or "page"


def to_full_url(url: str) -> str:
    return url if url.startswith("http") else f"{BASE}{url}"


def path_only(url: str) -> str:
    full = to_full_url(url)
    rest = full[len(BASE) :] if full.startswith(BASE) else url
    return rest or "/"


def goto(page, url: str) -> None:
    target = path_only(url)
    full = to_full_url(url)
    target_path = target.split("?")[0]

    try:
        link = page.locator(f'a[href="{target_path}"]')
        if link.count() > 0:
            link.first.click(force=True, timeout=8_000)
            page.wait_for_url(lambda u: path_only(u).split("?")[0] == target_path, timeout=12_000)
            return
    except PlaywrightError:
        pass

    label = PATH_CLICK_TEXT.get(target_path)
    if label:
        try:
            page.get_by_text(label, exact=True).first.click(timeout=5_000)
            page.wait_for_url(lambda u: path_only(u).split("?")[0] == target_path, timeout=12_000)
            return
        except PlaywrightError:
            pass

    try:
        page.goto(full, wait_until="commit", timeout=45_000)
        page.wait_for_url(lambda u: path_only(u).split("?")[0] == target_path, timeout=15_000)
    except PlaywrightError:
        pass
    try:
        page.wait_for_function(
            "() => document.querySelector('#root') && document.querySelector('#root').childElementCount > 0",
            timeout=12_000,
        )
    except PlaywrightTimeout:
        pass


def wait_settled(page, extra_ms: int = 800) -> None:
    try:
        page.wait_for_function(
            "() => document.body && document.body.innerText.trim().length > 20",
            timeout=12_000,
        )
    except PlaywrightTimeout:
        pass
    page.wait_for_timeout(extra_ms)
    try:
        page.locator("[aria-busy='true']").first.wait_for(state="hidden", timeout=2_500)
    except PlaywrightTimeout:
        pass


def login(page, role_path: str, username: str, password: str) -> None:
    page.goto(f"{BASE}{role_path}", wait_until="load", timeout=60_000)
    wait_settled(page, 600)
    page.locator('input[name="username"]').wait_for(timeout=20_000)
    page.locator('input[name="username"]').fill(username)
    page.locator('input[name="password"]').fill(password)
    page.get_by_role("button", name="Injira").click()
    page.wait_for_function(
        "() => !location.pathname.includes('/login/')",
        timeout=60_000,
    )
    wait_settled(page, 2500)


def screenshot_page(page, stem: str) -> list[Path]:
    SHOT_DIR.mkdir(parents=True, exist_ok=True)
    raw = SHOT_DIR / f"{stem}__full.jpg"
    kwargs = {
        "path": str(raw),
        "type": "jpeg",
        "quality": 82,
        "timeout": 15_000,
        "animations": "disabled",
    }
    # Viewport-only: full_page + ArcGIS WebGL has been crashing Chromium here.
    page.screenshot(full_page=False, **kwargs)
    return [raw]


def slice_image(raw: Path, stem: str) -> list[Path]:
    img = Image.open(raw)
    width, height = img.size
    if height <= SLICE_HEIGHT:
        return [raw]
    parts: list[Path] = []
    idx = 1
    top = 0
    while top < height:
        bottom = min(top + SLICE_HEIGHT, height)
        crop = img.crop((0, top, width, bottom))
        part = SHOT_DIR / f"{stem}__p{idx:02d}.jpg"
        crop.save(part, "JPEG", quality=78, optimize=True)
        parts.append(part)
        idx += 1
        top = bottom
    raw.unlink(missing_ok=True)
    return parts


def first_href(page, patterns: list[str]) -> str | None:
    try:
        hrefs = page.eval_on_selector_all(
            "a[href]",
            """els => els.map(e => e.getAttribute('href') || '')""",
        )
    except PlaywrightError:
        return None
    for href in hrefs:
        if not href or href.startswith("#"):
            continue
        for pattern in patterns:
            if re.search(pattern, href):
                # skip exact list roots
                if href.rstrip("/") in {
                    "/caretaker/abana",
                    "/district/abana",
                    "/district/ibigo",
                    "/district/abakoresha",
                    "/ncda/children",
                    "/ncda/centers",
                    "/ncda/users",
                    "/ncda/districts",
                    "/ncda/audit-logs",
                    "/caretaker/abakoresha",
                }:
                    continue
                return href
    return None


def capture(
    page,
    catalog: list[dict],
    section: str,
    title: str,
    path: str,
    extra_ms: int = 800,
    recover=None,
) -> None:
    if already(catalog, section, title):
        print(f"SKIP existing {section} / {title}", flush=True)
        return
    url = path if path.startswith("http") else f"{BASE}{path}"
    for attempt in range(2):
        try:
            if page.is_closed() and recover:
                page = recover()
            goto(page, url)
            wait_settled(page, extra_ms)
            n = len(catalog) + 1
            stem = f"{n:03d}-{slug(section)}-{slug(title)}"
            files = screenshot_page(page, stem)
            catalog.append(
                {
                    "section": section,
                    "title": title,
                    "route": path,
                    "final_url": page.url,
                    "files": [str(p) for p in files],
                }
            )
            print(f"OK  {section} / {title}  ({page.url})", flush=True)
            save_catalog(catalog)
            return
        except Exception as exc:  # noqa: BLE001
            closed = "closed" in str(exc).lower() or "Target" in str(exc)
            if closed and recover and attempt == 0:
                print(f"RECOVER after crash, retry {section} / {title}")
                page = recover()
                continue
            print(f"FAIL {section} / {title}  {path}  {exc}", flush=True)
            catalog.append(
                {
                    "section": section,
                    "title": f"{title} (capture failed)",
                    "route": path,
                    "final_url": path,
                    "files": [],
                    "error": str(exc),
                }
            )
            save_catalog(catalog)
            return


def click_first(page, names: list[str]) -> bool:
    for name in names:
        loc = page.get_by_role("button", name=re.compile(name))
        try:
            if loc.count() > 0:
                loc.first.click(timeout=5_000)
                wait_settled(page, 1000)
                return True
        except PlaywrightError:
            continue
        loc = page.get_by_role("link", name=re.compile(name))
        try:
            if loc.count() > 0:
                loc.first.click(timeout=5_000)
                wait_settled(page, 1000)
                return True
        except PlaywrightError:
            continue
    return False


def capture_if_href(
    page,
    catalog: list[dict],
    section: str,
    title: str,
    patterns: list[str],
    extra_ms: int = 1000,
    recover=None,
) -> str | None:
    href = first_href(page, patterns)
    if not href:
        print(f"SKIP detail {section} / {title} (no matching link)", flush=True)
        return None
    capture(page, catalog, section, title, href, extra_ms=extra_ms, recover=recover)
    return href


PUBLIC_PAGES = [
    ("Role selection", "/"),
    ("Login — Umurezi / ECD director", "/login/caretaker"),
    ("Login — Akarere", "/login/district"),
    ("Login — NCDA", "/login/ncda"),
    ("Forgot password", "/forgot-password"),
]

CARETAKER_PAGES = [
    ("Home", "/caretaker"),
    ("Children list", "/caretaker/abana"),
    ("Attendance", "/caretaker/ubwitabire"),
    ("Growth", "/caretaker/imikurire"),
    ("Monthly growth roster", "/caretaker/imikurire/ukwezi"),
    ("More (Ibindi)", "/caretaker/ibindi"),
    ("Register child", "/caretaker/kwiyandikisha"),
    ("Feeding (Imirire)", "/caretaker/imirire"),
    ("Feeding monthly report", "/caretaker/imirire/raporo"),
    ("STED", "/caretaker/sted"),
    ("STED wizard", "/caretaker/sted/new"),
    ("STED history", "/caretaker/sted/amateka"),
    ("Notifications", "/caretaker/amatangazo"),
    ("Impugukirwa", "/caretaker/impugukirwa"),
    ("Reports", "/caretaker/raporo"),
    ("Settings", "/caretaker/igenamiterere"),
]

DIRECTOR_PAGES = [
    ("Home", "/caretaker"),
    ("Center overview (Ikigo)", "/caretaker/ikigo"),
    ("Management hub", "/caretaker/imicungire"),
    ("Center users", "/caretaker/abakoresha"),
    ("Transfers", "/caretaker/kwimura"),
    ("Self-evaluation", "/caretaker/isuzuma"),
    ("Self-evaluation wizard", "/caretaker/isuzuma/new"),
    ("ECD Book hub", "/caretaker/igitabo"),
    ("Book VIII — Parent contributions", "/caretaker/igitabo/umusanzu"),
    ("Book IX — Parenting sessions", "/caretaker/igitabo/ibiganiro"),
    ("Book X — Committee", "/caretaker/igitabo/komite"),
    ("Book XI — Educators", "/caretaker/igitabo/abarezi"),
    ("Book XII — Support", "/caretaker/igitabo/ubufasha"),
    ("Book XIII — Visitors", "/caretaker/igitabo/abashyitsi"),
    ("Book XIV — Staff training", "/caretaker/igitabo/amahugurwa"),
    ("Children list", "/caretaker/abana"),
    ("Attendance", "/caretaker/ubwitabire"),
    ("Growth", "/caretaker/imikurire"),
    ("Notifications", "/caretaker/amatangazo"),
    ("Impugukirwa", "/caretaker/impugukirwa"),
    ("Settings", "/caretaker/igenamiterere"),
]

DISTRICT_PAGES = [
    ("Dashboard", "/district"),
    ("Centers", "/district/ibigo"),
    ("Children", "/district/abana"),
    ("Demographics", "/district/demografi"),
    ("Imikorere overview", "/district/imikorere"),
    ("Imikorere — Attendance", "/district/imikorere/ubwitabire"),
    ("Imikorere — Growth", "/district/imikorere/imikurire"),
    ("Imikorere — Feeding", "/district/imikorere/imirire"),
    ("Imikorere — STED", "/district/imikorere/sted"),
    ("Impugukirwa", "/district/impugukirwa"),
    ("Notifications", "/district/amatangazo"),
    ("Reports", "/district/raporo"),
    ("Referrals", "/district/referrals"),
    ("Map (Ikarita)", "/district/ikarita"),
    ("ECD Book hub", "/district/igitabo"),
    ("Book — Parent contributions", "/district/igitabo/umusanzu"),
    ("Book — Parenting sessions", "/district/igitabo/ibiganiro"),
    ("Book — Committee", "/district/igitabo/komite"),
    ("Book — Educators", "/district/igitabo/abarezi"),
    ("Book — Support", "/district/igitabo/ubufasha"),
    ("Book — Visitors", "/district/igitabo/abashyitsi"),
    ("Book — Staff training", "/district/igitabo/amahugurwa"),
    ("Caregivers", "/district/abakoresha"),
    ("Settings", "/district/igenamiterere"),
]

NCDA_PAGES = [
    ("Overview dashboard", "/ncda/dashboard"),
    ("Gukurikirana overview", "/ncda/gukurikirana"),
    ("Gukurikirana — Attendance", "/ncda/gukurikirana/ubwitabire"),
    ("Gukurikirana — Growth", "/ncda/gukurikirana/imikurire"),
    ("Gukurikirana — Feeding", "/ncda/gukurikirana/imirire"),
    ("Gukurikirana — STED", "/ncda/gukurikirana/sted"),
    ("Impugukirwa", "/ncda/gukurikirana/impugukirwa"),
    ("Inspections", "/ncda/inspections"),
    ("Children", "/ncda/children"),
    ("Demographics", "/ncda/demographics"),
    ("Notifications", "/ncda/amatangazo"),
    ("Users", "/ncda/users"),
    ("Roles", "/ncda/roles"),
    ("Settings", "/ncda/settings"),
    ("Audit logs", "/ncda/audit-logs"),
    ("Districts", "/ncda/districts"),
    ("Centers", "/ncda/centers"),
    ("ECD Book hub", "/ncda/igitabo"),
    ("Book — Parent contributions", "/ncda/igitabo/umusanzu"),
    ("Book — Parenting sessions", "/ncda/igitabo/ibiganiro"),
    ("Book — Committee", "/ncda/igitabo/komite"),
    ("Book — Educators", "/ncda/igitabo/abarezi"),
    ("Book — Support", "/ncda/igitabo/ubufasha"),
    ("Book — Visitors", "/ncda/igitabo/abashyitsi"),
    ("Book — Staff training", "/ncda/igitabo/amahugurwa"),
    ("WASH", "/ncda/wash"),
]


def set_heading_style(document: Document) -> None:
    styles = document.styles
    for name, size in (("Title", 28), ("Heading 1", 20), ("Heading 2", 14)):
        style = styles[name]
        style.font.color.rgb = RGBColor(0x1B, 0x3A, 0x4B)
        style.font.size = Pt(size)


def add_picture_fitted(doc: Document, path: str) -> None:
    pic = doc.add_picture(path, width=Inches(6.4))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    # Keep very tall slices from overflowing a page unreasonably.
    inline = pic._inline
    extent = inline.find(qn("wp:extent"))
    if extent is not None:
        cx = int(extent.get("cx"))
        cy = int(extent.get("cy"))
        max_cy = int(9.2 * 914400)
        if cy > max_cy and cx:
            scale = max_cy / cy
            extent.set("cx", str(int(cx * scale)))
            extent.set("cy", str(max_cy))


def build_docx(catalog: list[dict]) -> None:
    doc = Document()
    set_heading_style(doc)
    section = doc.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    doc.add_heading("ECD system UI catalog", 0)
    p = doc.add_paragraph(
        "Screenshots of Sisitemu y'Ubwitabire bw'Abana across public auth, "
        "Umurezi (caretaker), ECD director, Akarere (district), and NCDA admin. "
        "Captured from the local app on 8 September 2026. Passwords are not included."
    )
    p.runs[0].font.size = Pt(11)

    doc.add_paragraph(f"Source: {BASE}")
    doc.add_paragraph(f"Pages captured: {sum(1 for i in catalog if i.get('files'))}")

    doc.add_heading("Contents", 1)
    current = None
    for item in catalog:
        if item["section"] != current:
            current = item["section"]
            doc.add_paragraph(current, style="Heading 2")
        line = f"{item['title']} — {item['route']}"
        if not item.get("files"):
            line += " (not captured)"
        doc.add_paragraph(line, style="List Bullet")

    current = None
    for item in catalog:
        if item["section"] != current:
            current = item["section"]
            doc.add_heading(current, 1)
        doc.add_heading(item["title"], 2)
        meta = doc.add_paragraph()
        run = meta.add_run(f"Route: {item['route']}")
        run.italic = True
        run.font.size = Pt(10)
        run.font.color.rgb = RGBColor(0x4A, 0x55, 0x68)
        if item.get("error"):
            err = doc.add_paragraph(f"Capture error: {item['error']}")
            err.runs[0].font.color.rgb = RGBColor(0xB4, 0x23, 0x18)
        files = item.get("files") or []
        if len(files) > 1:
            doc.add_paragraph(
                f"Long page split into {len(files)} parts (top to bottom)."
            )
        for path in files:
            add_picture_fitted(doc, path)
            doc.add_paragraph("")

    DOCX_PATH.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(DOCX_PATH))
    print(f"Wrote {DOCX_PATH}", flush=True)


def run() -> None:
    SHOT_DIR.mkdir(parents=True, exist_ok=True)
    catalog: list[dict] = load_catalog()
    atexit.register(lambda: build_docx(load_catalog()))

    pw = sync_playwright().start()
    try:
        browser = pw.chromium.launch(
            headless=True,
            args=[
                "--disable-gpu",
                "--disable-dev-shm-usage",
                "--no-sandbox",
                "--use-angle=swiftshader",
                "--disable-software-rasterizer",
            ],
        )

        def recycle():
            nonlocal pw, browser
            try:
                browser.close()
            except Exception:
                pass
            try:
                pw.stop()
            except Exception:
                pass
            pw = sync_playwright().start()
            browser = pw.chromium.launch(
                headless=True,
                args=[
                    "--disable-gpu",
                    "--disable-dev-shm-usage",
                    "--no-sandbox",
                    "--use-angle=swiftshader",
                    "--disable-software-rasterizer",
                ],
            )

        def relaunch():
            nonlocal browser
            try:
                browser.close()
            except Exception:
                pass
            browser = pw.chromium.launch(
                headless=True,
                args=[
                    "--disable-gpu",
                    "--disable-dev-shm-usage",
                    "--no-sandbox",
                    "--use-angle=swiftshader",
                    "--disable-software-rasterizer",
                ],
            )
            return browser

        def new_page(width: int = 1280, height: int = 800):
            last = None
            for _ in range(2):
                try:
                    ctx = browser.new_context(
                        viewport={"width": width, "height": height},
                        ignore_https_errors=True,
                    )
                    pg = ctx.new_page()
                    pg.set_default_timeout(45_000)
                    pg.set_default_navigation_timeout(60_000)
                    return ctx, pg
                except Exception as exc:  # noqa: BLE001
                    last = exc
                    relaunch()
            raise last

        def make_session(login_path, user, password, width=1280, height=800):
            ctx, pg = new_page(width, height)
            login(pg, login_path, user, password)
            state = {"ctx": ctx, "page": pg}

            def recover():
                try:
                    state["ctx"].close()
                except Exception:
                    pass
                ctx2, pg2 = new_page(width, height)
                login(pg2, login_path, user, password)
                state["ctx"] = ctx2
                state["page"] = pg2
                return pg2

            return state, recover

        skip_public = all(
            already(catalog, "Public / authentication", title) for title, _ in PUBLIC_PAGES
        )
        if skip_public:
            print("SKIP public pages (already captured)", flush=True)
        else:
            context, page = new_page()
            for title, path in PUBLIC_PAGES:
                capture(page, catalog, "Public / authentication", title, path, extra_ms=500)
                if page.is_closed():
                    context, page = new_page()
            context.close()
        recycle()

        skip_caretaker = all(
            already(catalog, "Umurezi (caretaker)", title) for title, _ in CARETAKER_PAGES
        )
        if skip_caretaker:
            print("SKIP role Umurezi (already captured)", flush=True)
        else:
            user, password = CREDS["caretaker"]
            state, recover = make_session("/login/caretaker", user, password, 1100, 900)
        for title, path in ([] if skip_caretaker else CARETAKER_PAGES):
            extra = 1800 if "impugukirwa" in path else 800
            if path in {
                "/caretaker/kwiyandikisha",
                "/caretaker/imirire",
                "/caretaker/sted",
                "/caretaker/raporo",
                "/caretaker/igenamiterere",
            }:
                try:
                    goto(state["page"], "/caretaker/ibindi")
                    wait_settled(state["page"], 600)
                except PlaywrightError:
                    pass
            capture(
                state["page"],
                catalog,
                "Umurezi (caretaker)",
                title,
                path,
                extra_ms=extra,
                recover=recover,
            )
            if path == "/caretaker/abana":
                child_href = capture_if_href(
                    state["page"],
                    catalog,
                    "Umurezi (caretaker)",
                    "Child detail — overview",
                    [r"/caretaker/abana/[^/?#]+"],
                    recover=recover,
                )
                if not child_href:
                    try:
                        state["page"].get_by_text("Umwaka wa", exact=False).first.click(timeout=5_000)
                        wait_settled(state["page"], 1200)
                    except PlaywrightError:
                        pass
                    child_href = capture_if_href(
                        state["page"],
                        catalog,
                        "Umurezi (caretaker)",
                        "Child detail — overview",
                        [r"/caretaker/abana/[^/?#]+"],
                        recover=recover,
                    )
                if not child_href:
                    if click_first(state["page"], [r"^Reba ibirambuye", r"^Reba$"]):
                        n = len(catalog) + 1
                        stem = f"{n:03d}-umurezi-child-detail-overview"
                        files = screenshot_page(state["page"], stem)
                        catalog.append(
                            {
                                "section": "Umurezi (caretaker)",
                                "title": "Child detail — overview",
                                "route": path_only(state["page"].url),
                                "final_url": state["page"].url,
                                "files": [str(p) for p in files],
                            }
                        )
                        print(f"OK  Umurezi (caretaker) / Child detail — overview  ({state['page'].url})")
                        child_href = path_only(state["page"].url)
                if child_href:
                    base_child = child_href.split("?")[0]
                    for tab, label in (
                        ("profile", "Child detail — profile"),
                        ("attendance", "Child detail — attendance"),
                        ("growth", "Child detail — growth"),
                    ):
                        capture(
                            state["page"],
                            catalog,
                            "Umurezi (caretaker)",
                            label,
                            f"{base_child}?tab={tab}",
                            recover=recover,
                        )
                    capture(
                        state["page"],
                        catalog,
                        "Umurezi (caretaker)",
                        "Edit child",
                        f"{base_child}/hindura",
                        recover=recover,
                    )
        if not skip_caretaker:
            state["ctx"].close()
        recycle()

        skip_director = (
            sum(1 for title, _ in DIRECTOR_PAGES if already(catalog, "ECD director", title))
            >= max(1, len(DIRECTOR_PAGES) - 3)
        )
        if skip_director:
            print("SKIP role ECD director (already captured)", flush=True)
        else:
            try:
                user, password = CREDS["director"]
                state, recover = make_session("/login/caretaker", user, password, 1100, 900)
            except Exception as exc:  # noqa: BLE001
                print(f"LOGIN FAIL director: {exc}", flush=True)
                skip_director = True
        for title, path in ([] if skip_director else DIRECTOR_PAGES):
            capture(state["page"], catalog, "ECD director", title, path, recover=recover)
            if path == "/caretaker/abakoresha":
                capture_if_href(
                    state["page"],
                    catalog,
                    "ECD director",
                    "Center user detail",
                    [r"/caretaker/abakoresha/[^/?#]+"],
                    recover=recover,
                )
            if path == "/caretaker/abana":
                child_href = capture_if_href(
                    state["page"],
                    catalog,
                    "ECD director",
                    "Child detail",
                    [r"/caretaker/abana/[^/?#]+"],
                    recover=recover,
                )
                if not child_href:
                    click_first(state["page"], [r"^Reba ibirambuye", r"^Reba$"])
                    if "/abana/" in state["page"].url:
                        n = len(catalog) + 1
                        stem = f"{n:03d}-director-child-detail"
                        files = screenshot_page(state["page"], stem)
                        catalog.append(
                            {
                                "section": "ECD director",
                                "title": "Child detail",
                                "route": path_only(state["page"].url),
                                "final_url": state["page"].url,
                                "files": [str(p) for p in files],
                            }
                        )
                        print(f"OK  ECD director / Child detail  ({state['page'].url})")
        if not skip_director:
            state["ctx"].close()
        recycle()

        # District desktop
        try:
            user, password = CREDS["district"]
            state, recover = make_session("/login/district", user, password, 1440, 900)
        except Exception as exc:  # noqa: BLE001
            print(f"LOGIN FAIL district: {exc}", flush=True)
            recycle()
            user, password = CREDS["ncda"]
            state, recover = make_session("/login/ncda", user, password, 1440, 900)
            # fall through to NCDA by emptying district loop
            DISTRICT_PAGES_RUN = []
        else:
            DISTRICT_PAGES_RUN = list(DISTRICT_PAGES)
        for title, path in DISTRICT_PAGES_RUN:
            extra = 2500 if "ikarita" in path else 1200 if "impugukirwa" in path or path == "/district" else 800
            capture(
                state["page"],
                catalog,
                "Akarere (district)",
                title,
                path,
                extra_ms=extra,
                recover=recover,
            )
            if path == "/district/ibigo":
                capture_if_href(
                    state["page"],
                    catalog,
                    "Akarere (district)",
                    "Center detail",
                    [r"/district/ibigo/[^/?#]+"],
                    extra_ms=1500,
                    recover=recover,
                )
            if path == "/district/abana":
                href = capture_if_href(
                    state["page"],
                    catalog,
                    "Akarere (district)",
                    "Child detail",
                    [r"/district/abana/[^/?#]+"],
                    recover=recover,
                )
                if not href:
                    click_first(state["page"], [r"Reba Ibirambuye", r"^Reba$"])
            if path == "/district/abakoresha":
                capture_if_href(
                    state["page"],
                    catalog,
                    "Akarere (district)",
                    "Caregiver detail",
                    [r"/district/abakoresha/[^/?#]+"],
                    recover=recover,
                )
        state["ctx"].close()
        recycle()

        # NCDA desktop
        user, password = CREDS["ncda"]
        state, recover = make_session("/login/ncda", user, password, 1440, 900)
        for title, path in NCDA_PAGES:
            extra = 1500 if path in {"/ncda/dashboard", "/ncda/gukurikirana"} else 1000
            if "impugukirwa" in path:
                extra = 1500
            capture(
                state["page"],
                catalog,
                "NCDA admin",
                title,
                path,
                extra_ms=extra,
                recover=recover,
            )
            if path == "/ncda/children":
                href = capture_if_href(
                    state["page"],
                    catalog,
                    "NCDA admin",
                    "Child detail",
                    [r"/ncda/children/[^/?#]+"],
                    recover=recover,
                )
                if not href:
                    click_first(state["page"], [r"Reba ibirambuye", r"^Reba$"])
                    if "/children/" in state["page"].url:
                        n = len(catalog) + 1
                        stem = f"{n:03d}-ncda-child-detail"
                        files = screenshot_page(state["page"], stem)
                        catalog.append(
                            {
                                "section": "NCDA admin",
                                "title": "Child detail",
                                "route": path_only(state["page"].url),
                                "final_url": state["page"].url,
                                "files": [str(p) for p in files],
                            }
                        )
                        print(f"OK  NCDA admin / Child detail  ({state['page'].url})")
            if path == "/ncda/users":
                capture_if_href(
                    state["page"],
                    catalog,
                    "NCDA admin",
                    "User detail",
                    [r"/ncda/users/[^/?#]+"],
                    recover=recover,
                )
            if path == "/ncda/audit-logs":
                capture_if_href(
                    state["page"],
                    catalog,
                    "NCDA admin",
                    "Audit log detail",
                    [r"/ncda/audit-logs/[^/?#]+"],
                    recover=recover,
                )
            if path == "/ncda/districts":
                capture_if_href(
                    state["page"],
                    catalog,
                    "NCDA admin",
                    "District detail",
                    [r"/ncda/districts/[^/?#]+"],
                    recover=recover,
                )
            if path == "/ncda/centers":
                capture_if_href(
                    state["page"],
                    catalog,
                    "NCDA admin",
                    "Center detail",
                    [r"/ncda/centers/[^/?#]+"],
                    extra_ms=1500,
                    recover=recover,
                )
        state["ctx"].close()
        browser.close()
    finally:
        try:
            pw.stop()
        except Exception:
            pass

    save_catalog(catalog)
    build_docx(catalog)
    print(f"Catalog pages: {len(catalog)}", flush=True)


if __name__ == "__main__":
    started = time.time()
    run()
    print(f"Done in {time.time() - started:.0f}s", flush=True)
